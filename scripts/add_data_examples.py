#!/usr/bin/env python3
"""Add a "数据示例" column to the first sheet of data/data_format.xlsx.

Each of the 143 business-field rows gets one sample value taken from the
generated simulation database, matched through source_field_mapping.
"""
from __future__ import annotations

import argparse
import json
import sqlite3
from copy import copy
from pathlib import Path

import openpyxl

EXAMPLE_HEADER = "数据示例"
MAX_EXAMPLE_LENGTH = 150


def sample_value(
    connection: sqlite3.Connection, table: str, column: str
) -> str:
    row = connection.execute(
        f'SELECT "{column}" FROM "{table}" WHERE "{column}" IS NOT NULL LIMIT 1'
    ).fetchone()
    if row is None:
        return "（当前无故障仿真，该表为空，暂无示例）"
    value = row[0]
    if isinstance(value, float):
        value = round(value, 3)
    text = str(value)
    if column == "matrix_json":
        matrix = json.loads(text)
        return f"60×60邻接矩阵JSON：{json.dumps(matrix[0], separators=(',', ':'))[:100]}…（共{len(matrix)}行）"
    if len(text) > MAX_EXAMPLE_LENGTH:
        text = text[: MAX_EXAMPLE_LENGTH - 1] + "…"
    return text


def add_examples(workbook: Path, database: Path) -> int:
    connection = sqlite3.connect(f"file:{database}?mode=ro", uri=True)
    mapping = {}
    for source, field, table, column, modeled in connection.execute(
        """
        SELECT source_name, business_field_name, source_table, storage_column, modeled
        FROM source_field_mapping
        """
    ):
        mapping[(source, field)] = (table, column, modeled)

    wb = openpyxl.load_workbook(workbook)
    ws = wb[wb.sheetnames[0]]
    header_column = ws.max_column + 1
    header_cell = ws.cell(1, header_column, EXAMPLE_HEADER)
    header_cell._style = copy(ws.cell(1, ws.max_column - 1)._style)

    current_source = None
    written = 0
    for row_index in range(2, ws.max_row + 1):
        source = ws.cell(row_index, 1).value
        if source:
            current_source = str(source).strip()
        field = ws.cell(row_index, 2).value
        if field is None:
            continue
        field = str(field).strip()
        try:
            table, column, modeled = mapping[(current_source, field)]
        except KeyError:
            raise ValueError(f"No mapping for row {row_index}: {current_source} / {field}")
        if modeled and column is not None:
            example = sample_value(connection, table, column)
        else:
            example = "（按需求暂不仿真）"
        cell = ws.cell(row_index, header_column, example)
        cell._style = copy(ws.cell(row_index, 2)._style)
        cell.alignment = copy(ws.cell(row_index, 2).alignment)
        written += 1

    ws.column_dimensions[openpyxl.utils.get_column_letter(header_column)].width = 45
    wb.save(workbook)
    connection.close()
    return written


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=Path("data/data_format.xlsx"))
    parser.add_argument(
        "--database", type=Path, default=Path("data/sbc_simulation_20260808.db")
    )
    args = parser.parse_args()
    written = add_examples(args.workbook, args.database)
    print(f"{args.workbook}: wrote {written} examples")


if __name__ == "__main__":
    main()
